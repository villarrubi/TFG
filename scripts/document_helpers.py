"""Helpers compartidos por los scripts históricos de edición de documentos."""

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor, Twips
from docx.text.paragraph import Paragraph


def _paragraph(document, text):
    return next(p for p in document.paragraphs if p.text.strip() == text.strip())


def _set_text(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
    else:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""


def _replace_prefix(document, prefix, replacement):
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(prefix):
            _set_text(paragraph, replacement)


def _insert_paragraph_before(document, anchor, text, style=None):
    return anchor.insert_paragraph_before(text, style=style)


def _insert_paragraph_after(anchor, text, style=None):
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style or anchor.style
    paragraph.add_run(text)
    return paragraph


def _format_table(table, widths):
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths[index])
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string("FFFFFF")
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "004C3F")
                cell._tc.get_or_add_tcPr().append(shading)


def _insert_table_before(document, anchor, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for values in rows:
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
    anchor._p.addprevious(table._tbl)
    _format_table(table, widths)
    return table
