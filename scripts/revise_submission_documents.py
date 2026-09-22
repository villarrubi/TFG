"""Aplica la revisión final de memoria y guías solicitada por el tutor."""

from __future__ import annotations

from pathlib import Path

from document_helpers import (
    _format_table,
    _insert_paragraph_after,
    _insert_paragraph_before,
    _insert_table_before,
    _paragraph,
    _replace_prefix,
    _set_text,
)
from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]


def _insert_reference(document, anchor_prefix: str, text: str) -> None:
    author_prefix = text.split(" (", maxsplit=1)[0]
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(author_prefix)
        and paragraph.style.name == "Referencia TFG"
    ]
    if matches:
        reference = matches[0]
        _set_text(reference, text)
        for duplicate in matches[1:]:
            duplicate._element.getparent().remove(duplicate._element)
    else:
        anchor = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith(anchor_prefix)
        )
        reference = _insert_paragraph_before(
            document, anchor, text, style="Referencia TFG"
        )
    reference.paragraph_format.first_line_indent = Inches(-0.25)
    reference.paragraph_format.left_indent = Inches(0.25)


def _renumber_section_five_tables(document) -> None:
    for paragraph in document.paragraphs:
        if "Requisitos funcionales y no funcionales del prototipo." in paragraph.text:
            _set_text(paragraph, "Tabla 5.1: Requisitos funcionales y no funcionales del prototipo.")
        elif "Cambios respecto a la propuesta inicial." in paragraph.text:
            _set_text(paragraph, "Tabla 5.2: Cambios respecto a la propuesta inicial.")
        elif "Hiperparámetros finales del pipeline TF-IDF + MLP." in paragraph.text:
            _set_text(paragraph, "Tabla 5.3: Hiperparámetros finales del pipeline TF-IDF + MLP.")
        elif "Capas y responsabilidades del diseño modular." in paragraph.text:
            _set_text(paragraph, "Tabla 5.4: Capas y responsabilidades del diseño modular.")


def _insert_requirements(document) -> None:
    if any(
        paragraph.text == "Requisitos funcionales y no funcionales"
        for paragraph in document.paragraphs
    ):
        return
    anchor = _paragraph(document, "Cambios respecto a la propuesta inicial")
    _insert_paragraph_before(
        document,
        anchor,
        "Requisitos funcionales y no funcionales",
        style="Heading 2",
    )
    _insert_paragraph_before(
        document,
        anchor,
        "Los requisitos siguientes convierten el alcance descrito en criterios "
        "verificables. Los funcionales definen qué debe hacer el prototipo; los no "
        "funcionales fijan las condiciones de seguridad, reproducibilidad, "
        "mantenibilidad y rendimiento propias de una demostración académica local.",
    )
    caption = _insert_paragraph_before(
        document,
        anchor,
        "Tabla 5.1: Requisitos funcionales y no funcionales del prototipo.",
        style="Pie TFG",
    )
    caption.paragraph_format.keep_with_next = True
    _insert_table_before(
        document,
        anchor,
        ["Tipo", "ID", "Requisito", "Verificación"],
        [
            ["Funcional", "RF-01", "Analizar texto, EML y mensajes autorizados de Gmail.", "Pruebas de parser, API, Gmail y recorridos Chromium."],
            ["Funcional", "RF-02", "Ofrecer modos heurístico, neuronal y combinado.", "Evaluación de los tres modos sobre los mismos 16 EML."],
            ["Funcional", "RF-03", "Mostrar puntuación, veredicto, señales y explicación.", "Contrato JSON, interfaz y capturas reproducibles."],
            ["Funcional", "RF-04", "Mantener una versión activa ES y otra EN en el backend.", "Health, metadatos, activación atómica y pruebas."],
            ["Funcional", "RF-05", "Permitir alertas opcionales por Telegram.", "Prueba integral Gmail-HTTP-backend-Telegram con dobles."],
            ["No funcional", "RNF-01", "Procesar por defecto en loopback y minimizar datos persistidos.", "Rutas runtime separadas, OAuth de solo lectura y revisión de secretos."],
            ["No funcional", "RNF-02", "Poder reconstruir entrenamiento y evaluación.", "Dependencias fijadas, semillas, SHA-256 y scripts reproducibles."],
            ["No funcional", "RNF-03", "Separar presentación, transporte, aplicación y dominio.", "Arquitectura modular y suite automatizada."],
            ["No funcional", "RNF-04", "Arrancar y responder con latencia apta para la demo.", "Benchmark medido y prueba local de concurrencia."],
        ],
        [1050, 850, 3660, 2824],
    )
    _insert_paragraph_before(
        document,
        anchor,
        "El cumplimiento se evalúa dentro del alcance declarado: no implica eficacia "
        "universal, operación multiusuario ni garantías de producción.",
    )


def _insert_hyperparameters(document) -> None:
    if any(
        "Hiperparámetros finales del pipeline TF-IDF + MLP." in paragraph.text
        for paragraph in document.paragraphs
    ):
        return
    anchor = _paragraph(document, "Interfaz de usuario")
    _insert_paragraph_before(
        document,
        anchor,
        "La configuración final usa unigramas y bigramas, limita el vocabulario y "
        "mantiene una red compacta. La tabla recoge los valores efectivos de ambos "
        "modelos; las stopwords son las de scikit-learn para inglés y una lista local "
        "versionada para español.",
    )
    caption = _insert_paragraph_before(
        document,
        anchor,
        "Tabla 5.3: Hiperparámetros finales del pipeline TF-IDF + MLP.",
        style="Pie TFG",
    )
    caption.paragraph_format.keep_with_next = True
    _insert_table_before(
        document,
        anchor,
        ["Componente", "Parámetro", "Valor final"],
        [
            ["TF-IDF", "ngram_range", "(1, 2)"],
            ["TF-IDF", "max_features", "3.000"],
            ["TF-IDF", "min_df", "1"],
            ["TF-IDF", "strip_accents / stopwords", "unicode / inglés incorporadas o español local"],
            ["MLP", "hidden_layer_sizes", "(64, 32)"],
            ["MLP", "activation", "relu"],
            ["MLP", "alpha / learning_rate_init", "0,0001 / 0,001"],
            ["MLP", "max_iter / early_stopping", "500 / False"],
            ["MLP", "random_state", "42"],
        ],
        [1650, 2800, 3934],
    )
    _insert_paragraph_before(
        document,
        anchor,
        "Los artefactos activos se entrenaron el 30 de agosto de 2026: ES con 1.148 "
        "muestras (SHA-256 165e7c2bf292adf1d7bc88d936b3c14f4c7fe8f1caa3d2615718ca1190aaefb0) "
        "y EN con 65.661 (SHA-256 a3dd9dc3216445c70574982ad7b2515e02830e1a8c0f2ad841cf2c2eb2c56d69). "
        "Los hashes permiten comprobar exactamente qué ficheros se utilizaron en la entrega.",
    )


def _update_front_table_index(document) -> None:
    abbreviations = _paragraph(document, "Abreviaturas")
    front = []
    for paragraph in document.paragraphs:
        if paragraph._p is abbreviations._p:
            break
        if paragraph.text.startswith("Tabla "):
            front.append(paragraph)
    for text, before in (
        (
            "Tabla 5.1: Requisitos funcionales y no funcionales del prototipo.",
            "Tabla 5.2: Cambios respecto a la propuesta inicial.",
        ),
        (
            "Tabla 5.3: Hiperparámetros finales del pipeline TF-IDF + MLP.",
            "Tabla 5.4: Capas y responsabilidades del diseño modular.",
        ),
    ):
        if any(paragraph.text == text for paragraph in front):
            continue
        anchor = next(paragraph for paragraph in front if paragraph.text == before)
        _insert_paragraph_before(document, anchor, text)
    front = []
    for paragraph in document.paragraphs:
        if paragraph._p is abbreviations._p:
            break
        if paragraph.text.startswith("Tabla 5."):
            front.append(paragraph)
    hyperparameters = next(
        paragraph for paragraph in front if paragraph.text.startswith("Tabla 5.3:")
    )
    modular = next(
        paragraph for paragraph in front if paragraph.text.startswith("Tabla 5.4:")
    )
    modular._p.addprevious(hyperparameters._p)


def _update_experiment(document) -> None:
    _replace_prefix(
        document,
        "El protocolo reproducible fija las fuentes",
        "El protocolo reproducible fija fuentes, licencias y SHA-256 sin versionar "
        "mensajes. En español combina el entrenamiento de Softecapps (2024), con DOI "
        "10.57967/hf/2264, y SMS Spam Mexico (Iván, 2026): tras eliminar duplicados, "
        "conflictos y una coincidencia con el test quedan 1.148 textos, mientras 209 "
        "del split oficial se reservan para prueba. En inglés usa el agregado descrito "
        "por Al-Subaiey et al. (2024), que reúne CEAS, Enron, Ling, Nazario, Nigerian "
        "Fraud y SpamAssassin; no añade de nuevo sus componentes. Tras retirar 408 "
        "duplicados quedan 82.077 textos, divididos de forma estratificada 80/20 con "
        "semilla 42. Los joblib guardan el pipeline y metadatos, nunca el texto bruto.",
    )
    _replace_prefix(
        document,
        "La separación se realiza antes del ajuste.",
        "La separación se realiza antes del ajuste. Español conserva el test oficial "
        "sin coincidencias exactas e inglés deduplica antes de dividir. Otros 40 casos "
        "sintéticos y bilingües calibran la fusión mediante cinco particiones; 16 EML "
        "distintos prueban después los tres modos con MIME y cabeceras. La validación "
        "secundaria emplea una muestra única del conjunto de Miltchev et al. (2024) y "
        "DIFrauD; este último no se considera independiente por posible solapamiento de "
        "fuentes históricas con el agregado inglés (Boumber et al., 2024).",
    )
    _replace_prefix(
        document,
        "evaluation/training_sources.json fija",
        "evaluation/training_sources.json fija las URLs, licencias y huellas de los "
        "CSV; scripts/retrain_reproducible.py verifica, divide, entrena y evalúa. Los "
        "campos históricos llamados phishing/legitimate representan en realidad clase "
        "positiva (1) y negativa (0): los corpus ES son spam/ham usados como proxy de "
        "smishing/phishing y la clase positiva EN mezcla phishing y spam. No se equiparan "
        "semánticamente spam y phishing. Los holdouts de texto, además, carecen de MIME; "
        "estas limitaciones impiden extrapolar las cifras a producción.",
    )
    for table in document.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == "Conjunto":
            table.rows[0].cells[2].text = "Clase positiva (1)"
            table.rows[0].cells[3].text = "Clase negativa (0)"
            _format_table(table, [1200, 850, 1100, 1100, 2300, 1834])
            break
    _replace_prefix(
        document,
        "Tras fijar 45 % de peso heurístico",
        "La calibración recorre 40 casos separados del entrenamiento y de la evaluación "
        "final, en cinco particiones estratificadas por idioma y etiqueta. Prueba pesos "
        "heurísticos del 20 % al 50 % en pasos de 5 (el resto es neuronal), umbrales del "
        "20 al 60 en pasos de 1 y alta confianza del 65 al 85 en pasos de 5. Ordena por "
        "la menor accuracy balanceada entre particiones, su media y el valor global; "
        "después usa F1, recall y precisión. Los desempates predefinidos prefieren umbral "
        "cercano a 45, alta confianza cercana a 70 y mayor peso neuronal.",
    )
    anchor = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("La calibración recorre 40 casos separados")
    )
    combination_text = (
        "La combinación seleccionada fue 45 % heurístico, 55 % neuronal, umbral "
        "21 y alta confianza 70: obtuvo 0,625 de accuracy balanceada mínima entre "
        "particiones, 0,825 de media/global, F1 0,8293, recall 0,85 y precisión "
        "0,8095. El reparto 50/50 empató en métricas y perdió por el desempate "
        "declarado. El umbral 21 se optimizó para el combinado y se conserva como "
        "política común configurable; la puntuación resultante no es una probabilidad calibrada."
    )
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("La combinación seleccionada fue 45 %")
    ]
    if matches:
        _set_text(matches[0], combination_text)
        for duplicate in matches[1:]:
            duplicate._element.getparent().remove(duplicate._element)
    else:
        _insert_paragraph_after(
            anchor,
            combination_text,
        )
    _replace_prefix(
        document,
        "La mezcla ponderada original podía diluir",
        "La mezcla ponderada original podía diluir una evidencia fuerte. El protocolo "
        "anterior explica la selección reproducible de 45/55, umbral 21 y alta confianza "
        "70 sin usar los 16 EML finales. La alta confianza conserva una evidencia "
        "individual concluyente; los indicios BEC aislados pesan poco y solo la "
        "coincidencia de los tres recibe refuerzo. El resultado resuelve los escenarios "
        "BEC reservados, no todo BEC posible.",
    )


def _update_environment(document) -> None:
    if not any(
        paragraph.text == "Entorno de ejecución medido"
        for paragraph in document.paragraphs
    ):
        anchor = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith("El perfilado localizó")
        )
        _insert_paragraph_before(
            document, anchor, "Entorno de ejecución medido", style="Heading 3"
        )
        _insert_paragraph_before(
            document,
            anchor,
            "Las cifras se obtuvieron en Windows 11 Home 25H2 de 64 bits "
            "(compilación 26200.9278), con AMD Ryzen 7 7800X3D de 8 núcleos y 16 hilos "
            "y 63,1 GiB de RAM física. El entorno utilizó Python 3.12.7, scikit-learn "
            "1.9.0, Streamlit 1.58.0, NumPy 2.4.6, SciPy 1.17.1 y joblib 1.5.3. "
            "requirements.txt y constraints.txt fijan estas versiones para repetir la prueba.",
        )


def _update_terms_and_references(document) -> None:
    _replace_prefix(
        document,
        "SPF: Sender Policy Framework;",
        "SPF: Sender Policy Framework; política DNS que autoriza servidores para el "
        "dominio del sobre SMTP (MAIL FROM/HELO), no el campo From visible.",
    )
    _replace_prefix(
        document,
        "DMARC: Domain-based Message Authentication",
        "DMARC: Domain-based Message Authentication, Reporting and Conformance; "
        "política y mecanismo de informes basado en la alineación de SPF o DKIM con el dominio From visible.",
    )
    _replace_prefix(
        document,
        "DMARC: Política que indica",
        "DMARC: Política e informes del dominio del From visible. Exige alineación con "
        "al menos un SPF o DKIM válido y puede solicitar none, quarantine o reject; no "
        "verifica por sí mismo el contenido del mensaje.",
    )
    if not any(p.text.startswith("SPF: Política DNS") for p in document.paragraphs):
        dmarc = next(p for p in document.paragraphs if p.text.startswith("DMARC: Política e informes"))
        _insert_paragraph_before(
            document,
            dmarc,
            "SPF: Política DNS que permite comprobar si la IP emisora está autorizada "
            "para el dominio del sobre SMTP. No autentica por sí sola el From visible ni "
            "garantiza la integridad del contenido.",
        )
    _replace_prefix(
        document,
        "El código completo del prototipo se encuentra disponible",
        "El código completo del prototipo, las dependencias fijadas, los scripts de "
        "evaluación y las instrucciones de ejecución están disponibles en el repositorio "
        "público https://github.com/villarrubi/TFG.",
    )
    _replace_prefix(
        document,
        "Este capítulo reúne el material complementario del Trabajo Fin de Grado:",
        "Este capítulo reúne el material complementario del Trabajo Fin de Grado: "
        "la estructura completa del repositorio, los fragmentos de código fuente más "
        "representativos del sistema y una guía resumida de instalación y uso. El "
        "código completo, las dependencias fijadas, los scripts de evaluación y las "
        "instrucciones se encuentran en el repositorio público "
        "https://github.com/villarrubi/TFG.",
    )
    _replace_prefix(
        document,
        "Miltchev, R., Rangelov, N., & Genchev, K. (2024).",
        "Miltchev, R., Rangelov, D., & Genchev, E. (2024). Phishing validation "
        "emails dataset [Conjunto de datos]. Zenodo. "
        "https://doi.org/10.5281/zenodo.13474746",
    )
    _insert_reference(
        document,
        "Alkhalil, Z.",
        "Al-Subaiey, A., Al-Thani, M., Alam, N. A., Antora, K. F., Khandakar, A., & "
        "Zaman, S. A. U. (2024). Novel interpretable and robust web-based AI platform "
        "for phishing email detection. Computer & Electrical Engineering, 118, 109625. "
        "https://doi.org/10.1016/j.compeleceng.2024.109625 "
        "Dataset: https://www.kaggle.com/datasets/naserabdullahalam/phishing-email-dataset",
    )
    _insert_reference(
        document,
        "Kahneman, D.",
        "Iván, A. (2026). SMS Spam Mexico - Dataset en Español Mexicano [Conjunto de datos]. "
        "Kaggle. https://www.kaggle.com/datasets/aldoivan/sms-spam-mexico-dataset-en-espaol-mexicano",
    )
    _insert_reference(
        document,
        "Microsoft. (2025).",
        "Miltchev, R., Rangelov, D., & Genchev, E. (2024). Phishing validation emails dataset "
        "[Conjunto de datos]. Zenodo. https://doi.org/10.5281/zenodo.13474746",
    )
    _insert_reference(
        document,
        "Thakur, K.",
        "Softecapps. (2024). Spam/ham Spanish [Conjunto de datos]. Hugging Face. "
        "https://doi.org/10.57967/hf/2264",
    )


def revise_memory(path: Path) -> None:
    document = Document(path)
    _renumber_section_five_tables(document)
    _insert_requirements(document)
    _insert_hyperparameters(document)
    _update_front_table_index(document)
    _update_experiment(document)
    _update_environment(document)
    _update_terms_and_references(document)
    document.save(path)


def revise_full_guide(path: Path) -> None:
    document = Document(path)
    _replace_prefix(
        document,
        "Antes de entrenar, el cliente envía los CSV",
        "Antes de entrenar, el cliente envía los CSV a /datasets/summary. El backend "
        "valida columnas y devuelve filas, clase positiva (1) y clase negativa (0) por "
        "archivo. Esos nombres evitan equiparar phishing con spam en los corpus proxy; "
        "el navegador solo presenta el resumen.",
    )
    _replace_prefix(
        document,
        "Estas cifras son comparativas del entorno medido",
        "Estas cifras se midieron en Windows 11 Home 25H2 de 64 bits "
        "(compilación 26200.9278), AMD Ryzen 7 7800X3D (8 núcleos/16 hilos), 63,1 GiB "
        "de RAM y Python 3.12.7, con scikit-learn 1.9.0, Streamlit 1.58.0, NumPy 2.4.6, "
        "SciPy 1.17.1 y joblib 1.5.3. Son comparativas locales, no una promesa universal.",
    )
    _replace_prefix(
        document,
        "Los modelos activos declaran",
        "Los modelos activos declaran 1.148 muestras ES (613 positivas spam/phishing "
        "proxy y 535 negativas) y 65.661 EN (34.275 positivas phishing/spam y 31.386 "
        "negativas). No se equiparan semánticamente spam y phishing. Los hashes de los "
        "joblib son ES 165e7c2bf292 y EN a3dd9dc32164. Otros 40 casos calibran y 16 EML "
        "finales comparan los tres modos; el evaluador aporta métricas, matrices, hashes "
        "y detalle por caso con límites de representatividad documentados.",
    )
    _replace_prefix(
        document,
        "El valor 21 fue seleccionado",
        "El umbral 21 fue seleccionado sobre 40 casos separados, con cinco particiones "
        "estratificadas por idioma y etiqueta. La rejilla probó umbrales 20-60 y pesos "
        "heurísticos 20-50; priorizó la peor accuracy balanceada, su media/global, F1, "
        "recall y precisión. Es una política operativa del combinado, no una probabilidad, "
        "y debe confirmarse con correo real independiente.",
    )
    _replace_prefix(
        document,
        "La rejilla de calibración sobre 40 casos seleccionó",
        "La rejilla sobre 40 casos seleccionó 45 % heurístico y 55 % neuronal, umbral "
        "21 y alta confianza 70. El 50/50 empató en métricas y el desempate predefinido "
        "prefirió mayor peso neuronal; para alta confianza se eligió el valor más cercano "
        "a 70. El resultado fue accuracy balanceada mínima 0,625, media/global 0,825, "
        "F1 0,8293, recall 0,85 y precisión 0,8095. Los 16 EML se evaluaron después, sin reajustar.",
    )
    demo_steps = [
        "1.\tArranca backend y Streamlit desde una copia limpia; comprueba /health y las versiones ES/EN.",
        "2.\tAnaliza primero el EML legítimo en modo combinado y explica por qué queda por debajo del umbral.",
        "3.\tAnaliza el EML BEC/phishing en modo combinado y abre puntuación, señales y explicación.",
        "4.\tRepite el caso de phishing en modo heurístico para auditar las reglas activadas.",
        "5.\tRepite en modo neuronal y compara qué evidencia aporta y qué no puede observar sin cabeceras.",
        "6.\tAbre Entrenamiento > Modelos guardados y enseña idioma, muestras, hashes e hiperparámetros; no reentrenes en directo.",
        "7.\tMuestra el repositorio, los comandos reproducibles y la salida de 94 pruebas Python más 2 recorridos Chromium.",
        "8.\tCierra con límites: Gmail y Telegram son integraciones opcionales y solo se enseñan si las credenciales de prueba ya funcionan.",
    ]
    for number, replacement in enumerate(demo_steps, 1):
        _replace_prefix(document, f"{number}.\t", replacement)
    _replace_prefix(
        document,
        "SPF: autoriza",
        "SPF: autoriza por DNS los servidores del dominio del sobre SMTP; no autentica por sí solo el From visible.",
    )
    _replace_prefix(
        document,
        "DMARC: aplica",
        "DMARC: exige que SPF o DKIM válidos se alineen con el dominio From visible y publica política/informes.",
    )
    document.save(path)


def revise_short_guide(path: Path) -> None:
    document = Document(path)
    demo = [
        "Arrancar backend y Streamlit desde una copia limpia; comprobar /health y las versiones ES/EN.",
        "Analizar primero el EML legítimo en modo combinado y explicar el resultado bajo el umbral.",
        "Analizar el EML BEC/phishing en modo combinado y abrir señales, URLs y explicación.",
        "Repetir el phishing en modo heurístico y neuronal para comparar los tres modos sobre el mismo caso.",
        "Abrir Modelos guardados y mostrar idioma, muestras, hashes e hiperparámetros sin reentrenar.",
        "Mostrar el repositorio, los comandos reproducibles y la salida de 94 pruebas Python más 2 recorridos Chromium.",
        "Cerrar con los límites; Gmail y Telegram solo se enseñan si las credenciales de prueba ya están preparadas.",
    ]
    list_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "List Number"
    ]
    for paragraph, replacement in zip(list_paragraphs[:7], demo):
        _set_text(paragraph, replacement)
    _replace_prefix(
        document,
        "La calibración usa 40 casos distintos",
        "La calibración usa 40 casos distintos de los 16 EML finales y cinco "
        "particiones estratificadas. La rejilla prueba pesos heurísticos 20-50, umbrales "
        "20-60 y alta confianza 65-85; prioriza peor/media/global accuracy balanceada, "
        "F1, recall y precisión. Selecciona 45/55, umbral 21 y alta confianza 70; 50/50 "
        "empata y pierde por el desempate declarado. Sobre calibración logra 0,625 de "
        "accuracy balanceada mínima y 0,825 media/global. La puntuación no es una probabilidad.",
    )
    _replace_prefix(
        document,
        "En rendimiento, presenta solo mejoras medidas",
        "En rendimiento, presenta solo mejoras medidas en Windows 11 Home 25H2, Ryzen "
        "7 7800X3D, 63,1 GiB de RAM y Python 3.12.7: la importación fría de heurísticas "
        "bajó un 96,6 % y el arranque un 76,2 %. La inferencia varió menos de un 3 %. "
        "A 8 y 16 análisis simultáneos no hubo fallos en las rondas locales; a 32 sí, "
        "por lo que 4-8 es la recomendación de demo, no un SLA.",
    )
    document.save(path)


def revise_technology_guide(path: Path) -> None:
    document = Document(path)
    _replace_prefix(
        document,
        "Los hiperparámetros se concentran",
        "Los hiperparámetros finales son TF-IDF (ngram_range 1-2, max_features 3.000, "
        "min_df 1, strip_accents unicode) y MLP ((64, 32), relu, alpha 0,0001, learning "
        "rate 0,001, 500 iteraciones, sin early stopping, semilla 42). El protocolo usa "
        "1.148/209 textos ES y 65.661/16.416 EN. La clase positiva mezcla spam/phishing "
        "según la fuente y se documenta como proxy, no como equivalencia. Los hashes "
        "permiten identificar los joblib y los modelos no guardan textos brutos.",
    )
    _replace_prefix(
        document,
        "El benchmark reproducible separa",
        "El benchmark reproducible separa arranque e inferencia. Se midió en Windows 11 "
        "Home 25H2 de 64 bits (26200.9278), Ryzen 7 7800X3D, 63,1 GiB RAM y Python "
        "3.12.7 con scikit-learn 1.9.0, Streamlit 1.58.0, NumPy 2.4.6, SciPy 1.17.1 y "
        "joblib 1.5.3. La carga diferida redujo importación heurística un 96,6 % y "
        "arranque de aplicación un 76,2 %; la inferencia varió menos de ±3 %.",
    )
    _replace_prefix(
        document,
        "La prueba de concurrencia usa",
        "La calibración usa 40 casos separados en cinco particiones: pesos heurísticos "
        "20-50, umbrales 20-60 y alta confianza 65-85. Selecciona 45/55, umbral 21 y "
        "alta confianza 70 con desempates predefinidos; no produce una probabilidad "
        "calibrada. La prueba de concurrencia usa el backend real: 8 y 16 análisis "
        "terminaron sin errores, pero a 32 aparecieron fallos. Se recomiendan 4-8 para demo.",
    )
    document.save(path)


def main() -> None:
    revise_memory(ROOT / "TFG.docx")
    revise_full_guide(ROOT / "Guia_defensa_TFG.docx")
    revise_short_guide(ROOT / "Guia_03_Guion_defensa.docx")
    revise_technology_guide(ROOT / "Guia_02_Tecnologias_y_decisiones.docx")
    print("Memoria y tres guías revisadas.")


if __name__ == "__main__":
    main()
